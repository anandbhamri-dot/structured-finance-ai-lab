"""Generate a synthetic residential mortgage loan tape using standard field names.

Writes two files:
  sample_data/sample_loan_tape.csv   - plain values, for tape_analyser.py
  sample_data/sample_loan_tape.xlsx  - the same tape, with LOAN_AGE and CUR_LTV as live
                                        Excel formulas (field-level calculations done in Excel)

All data is randomly generated. A handful of deliberate data errors are injected so the
quality checks have something to find, as they would on a real originator tape.
"""
import argparse
from datetime import date

import numpy as np
import pandas as pd

STATES = ["CA", "TX", "FL", "NY", "IL", "WA", "AZ", "CO", "NC", "GA", "NJ", "VA", "MA", "OH", "PA"]
STATE_WEIGHTS = [0.22, 0.11, 0.10, 0.07, 0.05, 0.06, 0.05, 0.05, 0.05, 0.05, 0.04, 0.04, 0.04, 0.035, 0.035]
CUTOFF = date(2026, 9, 1)


def make_tape(n: int, seed: int) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    fico = np.clip(rng.normal(748, 38, n), 580, 830).round()
    ltv = np.clip(rng.normal(74, 12, n), 20, 97).round(1)
    cltv = np.clip(ltv + rng.choice([0, 0, 0, 5, 10], n), ltv, 100).round(1)
    dti = np.clip(rng.normal(36, 8, n), 8, 55).round(1)
    orig_bal = (np.clip(rng.lognormal(12.9, 0.45, n), 60_000, 3_000_000) / 1000).round() * 1000
    term = rng.choice([360, 360, 360, 180, 240], n)
    age = rng.integers(0, 36, n)
    rate = np.clip(6.9 - (fico - 748) * 0.006 + (ltv - 74) * 0.01 + rng.normal(0, 0.25, n), 4.5, 10.5).round(3)
    orig_dates = [date(CUTOFF.year + (CUTOFF.month - 1 - int(a)) // 12, (CUTOFF.month - 1 - int(a)) % 12 + 1, 1) for a in age]
    df = pd.DataFrame({
        "LOAN_ID": [f"LN{100000 + i}" for i in range(n)],
        "STATE": rng.choice(STATES, n, p=STATE_WEIGHTS),
        "PROP_TYPE": rng.choice(["SFR", "PUD", "Condo", "2-4 Family"], n, p=[0.62, 0.2, 0.12, 0.06]),
        "OCCUPANCY": rng.choice(["Owner", "Second Home", "Investor"], n, p=[0.86, 0.05, 0.09]),
        "PURPOSE": rng.choice(["Purchase", "Rate/Term Refi", "Cash-Out Refi"], n, p=[0.58, 0.22, 0.20]),
        "DOC_TYPE": rng.choice(["Full", "Bank Statement", "DSCR"], n, p=[0.8, 0.12, 0.08]),
        "FICO": fico.astype(int),
        "LTV": ltv,
        "CLTV": cltv,
        "DTI": dti,
        "ORIG_BAL": orig_bal,
        "PRIN_BAL": (orig_bal * (1 - age * 0.0025)).round(2),
        "PROP_VALUE": (orig_bal / (ltv / 100) / 1000).round() * 1000,
        "NOTE_RATE": rate,
        "ORIG_TERM": term,
        "REM_TERM": term - age,
        "LOAN_AGE": age,
        "ORIG_DATE": pd.to_datetime(orig_dates),
        "DQ_STATUS": rng.choice(["Current", "30", "60", "90+"], n, p=[0.955, 0.03, 0.01, 0.005]),
    })
    df["CUR_LTV"] = (df["PRIN_BAL"] / df["PROP_VALUE"] * 100).round(2)
    # Inject realistic tape errors
    idx = rng.choice(n, 12, replace=False)
    df.loc[idx[0:3], "FICO"] = 0                                        # missing FICO coded as zero
    df.loc[idx[3:5], "CLTV"] = df.loc[idx[3:5], "LTV"] - 8              # CLTV below LTV
    df.loc[idx[5:7], "PRIN_BAL"] = df.loc[idx[5:7], "ORIG_BAL"] * 1.2   # balance above original
    df.loc[idx[7], "STATE"] = None                                      # missing state
    df.loc[idx[8:10], "REM_TERM"] = df.loc[idx[8:10], "ORIG_TERM"] + 12 # remaining > original
    df.loc[idx[10], "NOTE_RATE"] = 69.5                                 # rate keyed without decimal
    df = pd.concat([df, df.iloc[[idx[11]]]], ignore_index=True)        # duplicate loan ID
    df["LOAN_AGE"] = df["ORIG_TERM"] - df["REM_TERM"]
    df["CUR_LTV"] = (df["PRIN_BAL"] / df["PROP_VALUE"] * 100).round(2)
    return df


def write_xlsx(df: pd.DataFrame, path: str) -> None:
    """Write the tape with LOAN_AGE and CUR_LTV as Excel formulas."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = list(df.columns)
    col = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
    wb = Workbook()
    ws = wb.active
    ws.title = "Tape"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F3864")
    for r, row in enumerate(df.itertuples(index=False), start=2):
        values = []
        for c, v in zip(cols, row):
            if c == "LOAN_AGE":
                v = f"={col['ORIG_TERM']}{r}-{col['REM_TERM']}{r}"
            elif c == "CUR_LTV":
                v = f"=IFERROR(ROUND({col['PRIN_BAL']}{r}/{col['PROP_VALUE']}{r}*100,2),\"\")"
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime().date()
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = None if np.isnan(v) else float(v)
            values.append(v)
        ws.append(values)
    for c in cols:
        ws.column_dimensions[col[c]].width = max(11, len(c) + 3)
        if c == "ORIG_DATE":
            for cell in ws[col[c]][1:]:
                cell.number_format = "yyyy-mm-dd"
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = Font(name="Arial")
    ws.freeze_panes = "B2"

    notes = wb.create_sheet("Notes")
    for line in [
        ["Synthetic sample loan tape (no real borrower data)"],
        [],
        ["LOAN_AGE and CUR_LTV are Excel formulas, showing field-level calculations done in Excel before loading:"],
        [f"LOAN_AGE = ORIG_TERM - REM_TERM  (e.g. ={col['ORIG_TERM']}2-{col['REM_TERM']}2)"],
        [f"CUR_LTV = PRIN_BAL / PROP_VALUE * 100  (e.g. =ROUND({col['PRIN_BAL']}2/{col['PROP_VALUE']}2*100,2))"],
        [],
        ["12 data errors are planted for the quality checks: zero FICOs, CLTV below LTV, balances above original,"],
        ["a missing state, remaining term above original, a rate keyed as 69.5 and a duplicate loan ID."],
    ]:
        notes.append(line)
    notes.column_dimensions["A"].width = 110
    for r in notes.iter_rows():
        for cell in r:
            cell.font = Font(name="Arial", bold=cell.row == 1)
    wb.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--loans", type=int, default=2500)
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--out", default="sample_data/sample_loan_tape")
    args = ap.parse_args()
    tape = make_tape(args.loans, args.seed)
    tape.to_csv(args.out + ".csv", index=False, date_format="%Y-%m-%d")
    write_xlsx(tape, args.out + ".xlsx")
    print(f"Wrote {len(tape):,} loans to {args.out}.csv and {args.out}.xlsx")
